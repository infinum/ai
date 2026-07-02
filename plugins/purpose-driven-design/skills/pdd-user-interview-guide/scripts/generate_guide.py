#!/usr/bin/env python3
"""
generate_guide.py — build a branded Infinum/ETR interview-guide .docx from a Markdown file.

Pure Python (python-docx) so it runs in the Cowork base image — no Node required, and each
skill ships its own copy (no cross-skill /tmp dependency).

Usage:
  python3 generate_guide.py --content guide.md --output "Guide.docx" [--logo path/to/logo-color.png]

Markdown conventions:
  # Title | Client | Date   → title block + red rule + grey subtitle
  ## Section                → H2 heading (also restarts question numbering)
  ### Subsection            → H3 heading
  1. text                   → numbered question (restarts at each ## section)
  a. text                   → lettered sub-item (probe)
  - text                    → bullet
  *italic*                  → italic note (grey);  **bold** markers render as normal weight
Every page carries the Infinum logo top-right (or the "INFINUM" wordmark if the asset is missing).
"""
import argparse, os, re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Helvetica Neue"
RED = RGBColor(0xD8, 0x26, 0x2D)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x77, 0x77, 0x77)
SECTION_GREY = RGBColor(0x55, 0x55, 0x55)


def find_logo(logo_arg):
    here = os.path.dirname(os.path.abspath(__file__))
    for c in [logo_arg,
              os.path.join(here, "..", "assets", "logo-color.png"),
              os.path.join(here, "..", "..", "..", "assets", "logo-color.png")]:
        if c and os.path.exists(c):
            return c
    return None


def set_bottom_border(paragraph, color="D8262D", sz="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", sz), ("w:space", "2"), ("w:color", color)):
        bottom.set(qn(k), v)
    pbdr.append(bottom)
    # w:pBdr must precede shd/spacing/ind/etc. in the CT_PPr child order — insert, don't append.
    pPr.insert_element_before(
        pbdr,
        "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
        "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
        "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
        "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
        "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
    )


def style_run(run, size, color=DARK, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def add_inline(paragraph, text, size=11, color=DARK):
    regex = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*")
    i = 0
    for m in regex.finditer(text):
        if m.start() > i:
            style_run(paragraph.add_run(text[i:m.start()]), size, color)
        if m.group(1) is not None:
            style_run(paragraph.add_run(m.group(1)), size, color)     # ** ** -> normal weight
        else:
            style_run(paragraph.add_run(m.group(2)), size, GREY, italic=True)  # * * -> italic grey
        i = m.end()
    if i < len(text):
        style_run(paragraph.add_run(text[i:]), size, color)
    if not paragraph.runs:
        style_run(paragraph.add_run(text), size, color)


def build(content_path, output_path, logo_path):
    md = open(content_path, encoding="utf-8").read()
    doc = Document()
    # python-docx ships a <w:zoom/> with no percent, which strict validators reject — set it.
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is None:
        zoom = OxmlElement("w:zoom"); doc.settings.element.insert(0, zoom)
    zoom.set(qn("w:percent"), "100")
    sec = doc.sections[0]
    sec.top_margin = Inches(1.0); sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.75); sec.right_margin = Inches(0.75)

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if logo_path:
        hp.add_run().add_picture(logo_path, width=Inches(1.6))
    else:
        style_run(hp.add_run("INFINUM"), 9, RED, bold=True)

    num = 0
    for raw in md.splitlines():
        s = raw.strip()
        if not s:
            continue
        if s.startswith("# ") and not s.startswith("## "):
            parts = [p.strip() for p in s[2:].split("|")]
            t = doc.add_paragraph(); style_run(t.add_run(parts[0]), 24, DARK, bold=True)
            rule = doc.add_paragraph(); set_bottom_border(rule); rule.paragraph_format.space_after = Pt(8)
            if len(parts) > 1:
                sub = doc.add_paragraph(); style_run(sub.add_run(" | ".join(parts[1:])), 10, GREY)
            continue
        if s.startswith("## ") and not s.startswith("### "):
            num = 0
            h = doc.add_paragraph(); h.paragraph_format.space_before = Pt(14)
            style_run(h.add_run(s[3:].strip()), 15, DARK, bold=True)
            continue
        if s.startswith("### "):
            h = doc.add_paragraph(); h.paragraph_format.space_before = Pt(10)
            style_run(h.add_run(s[4:].strip()), 11, SECTION_GREY, bold=True)
            set_bottom_border(h, color="DDDDDD", sz="4")
            continue
        m = re.match(r"^(\d+)\.\s+(.*)", s)
        if m:
            num += 1
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(18); p.paragraph_format.space_after = Pt(4)
            style_run(p.add_run(f"{num}. "), 11, DARK)
            add_inline(p, m.group(2))
            continue
        m = re.match(r"^([a-z])\.\s+(.*)", s)
        if m:
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(42)
            style_run(p.add_run(f"{m.group(1)}. "), 11, GREY)
            add_inline(p, m.group(2), color=GREY)
            continue
        m = re.match(r"^[-*]\s+(.*)", s)
        if m:
            p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.left_indent = Pt(24)
            add_inline(p, m.group(1))
            continue
        p = doc.add_paragraph(); add_inline(p, s)

    doc.save(output_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--content", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--logo", default=None)
    a = ap.parse_args()
    if not os.path.exists(a.content):
        sys.exit(f"Content file not found: {a.content}")
    logo = find_logo(a.logo)
    build(a.content, a.output, logo)
    print(f"Saved: {a.output}" + ("" if logo else "  (logo asset not found - used wordmark fallback)"))


if __name__ == "__main__":
    main()
