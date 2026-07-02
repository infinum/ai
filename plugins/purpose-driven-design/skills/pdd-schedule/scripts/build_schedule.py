#!/usr/bin/env python3
"""
build_schedule.py — generate a Project Schedule (Gantt) that matches the ETR/Infinum
"Project Schedule" tab. It edits a copy of the bundled template (so fonts, colours and the
week/day header styling are preserved), keeps only the Project Schedule sheet, and regenerates
the schedule from a config: a start date + phases, each with tasks and milestones.

Usage:
  python3 build_schedule.py --config config.json --output "Project Schedule - [Project].xlsx"
  (uses the bundled assets/schedule-template.xlsx unless --template is given)

config.json schema:
{
  "title": "Acme — Website Redesign",        # optional; written top-left
  "start": "2026-07-06",                      # ideally a Monday
  "weeks": 16,                                # optional; inferred from task dates if omitted
  "phases": [
    {"name": "Discovery Process", "color": "E1EC9E", "tasks": [
      {"num": "1.0", "name": "Discovery kickoff",          "start": "2026-07-08", "end": "2026-07-08", "milestone": true},
      {"num": "1.1", "name": "Stakeholder interviews",     "start": "2026-07-09", "end": "2026-07-20"},
      {"num": "1.2", "name": "Discovery readout (delivery)","start": "2026-07-24", "end": "2026-07-24", "milestone": true},
      {"num": "1.3", "name": "Client feedback due",        "start": "2026-07-29", "end": "2026-07-29", "milestone": true}
    ]},
    {"name": "User Experience", "color": "F6A3C0", "tasks": [ ... ]},
    {"name": "Visual Design",   "color": "9FB4DB", "tasks": [ ... ]},
    {"name": "Development",      "color": "B6D7A8", "tasks": [ ... ]}
  ]
}
Set "milestone": true on delivery dates and client-feedback-due dates (single-day markers,
bolded). Each phase gets its own colour and a heading row.
"""
import argparse, json, os, sys, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BUNDLED = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "..", "assets", "schedule-template.xlsx")
SCHEDULE_SHEET = "Project Schedule"
DEFAULT_PHASE_COLORS = ["E1EC9E", "F6A3C0", "9FB4DB", "B6D7A8", "C3D83B", "6E8EC8"]
DATE_COL_START = 5      # column E
DATE_ROW = 4
DOW_ROW = 3
WEEK_ROW = 1
FIRST_ROW = 6           # first phase/task row
DOW = {0: "M", 1: "T", 2: "W", 3: "Th", 4: "F", 5: "Sat", 6: "Sun"}


def pdate(s):
    if isinstance(s, (datetime.datetime, datetime.date)):
        return s.date() if isinstance(s, datetime.datetime) else s
    return datetime.datetime.strptime(str(s)[:10], "%Y-%m-%d").date()


def drange(s, e):
    s, e = pdate(s), pdate(e)
    return f"{s.month}/{s.day}" if s == e else f"{s.month}/{s.day}-{e.month}/{e.day}"


def build(ws, cfg):
    start = pdate(cfg["start"])
    phases = cfg.get("phases", [])
    weeks = cfg.get("weeks")
    if not weeks:
        last = start
        for ph in phases:
            for t in ph.get("tasks", []):
                last = max(last, pdate(t.get("end", t["start"])))
        weeks = max(1, ((last - start).days // 7) + 2)
    days = weeks * 7
    last_col = DATE_COL_START + days - 1

    for mr in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(mr))

    # date header (rows 1/3/4)
    for i in range(days):
        col = DATE_COL_START + i
        d = start + datetime.timedelta(days=i)
        ws.cell(DATE_ROW, col).value = datetime.datetime(d.year, d.month, d.day)
        ws.cell(DOW_ROW, col).value = DOW[d.weekday()]
        ws.cell(WEEK_ROW, col).value = (f"Week {i // 7 + 1}" + ("*" if i == 0 else "")) if i % 7 == 0 else None
    for col in range(last_col + 1, ws.max_column + 1):
        for r in (WEEK_ROW, DOW_ROW, DATE_ROW):
            ws.cell(r, col).value = None
    for w in range(weeks):
        c0 = DATE_COL_START + w * 7
        ws.merge_cells(start_row=WEEK_ROW, start_column=c0, end_row=WEEK_ROW, end_column=c0 + 6)

    # clear old content AND all fills below the header (the template's example left-column
    # bands + bars must be wiped, or they linger and no longer line up with the new rows)
    no_fill = PatternFill(fill_type=None)
    for r in range(FIRST_ROW, ws.max_row + 1):
        for c in range(1, 5):
            ws.cell(r, c).value = None
        for c in range(1, ws.max_column + 1):
            ws.cell(r, c).fill = no_fill

    # optional title (top-left, above the grid)
    if cfg.get("title"):
        t = ws.cell(2, 1, cfg["title"])
        t.font = Font(name="IBM Plex Serif", size=16, bold=True, color="000000")

    phase_font = Font(name="IBM Plex Sans", size=11, bold=True, color="FFFFFF")  # white on black band
    body = Font(name="Arial", size=10, color="000000")
    body_b = Font(name="Arial", size=10, bold=True, color="000000")
    left = Alignment(horizontal="left", vertical="center")
    dark_band = PatternFill("solid", fgColor="FF242526")  # heading marker in col A
    row = FIRST_ROW
    for idx, ph in enumerate(phases):
        color = ph.get("color") or DEFAULT_PHASE_COLORS[idx % len(DEFAULT_PHASE_COLORS)]
        if len(color) == 6:
            color = "FF" + color
        band = PatternFill("solid", fgColor=color)
        # phase heading row: black band across cols A–C with white heading text
        hc = ws.cell(row, 1, ph["name"]); hc.font = phase_font; hc.alignment = left
        for cc in (1, 2, 3):
            ws.cell(row, cc).fill = dark_band
        row += 1
        for t in ph.get("tasks", []):
            is_ms = bool(t.get("milestone"))
            ws.cell(row, 1).fill = band           # col A: phase band (lines up with this phase)
            if t.get("num") is not None:
                bn = ws.cell(row, 2, t["num"]); bn.font = body; bn.fill = band  # tinted number cell
            else:
                ws.cell(row, 2).fill = band
            name = ("◆ " + t["name"]) if is_ms else t["name"]
            cn = ws.cell(row, 3, name); cn.font = body_b if is_ms else body; cn.alignment = left
            ws.cell(row, 4, t.get("range") or drange(pdate(t["start"]), pdate(t.get("end", t["start"])))).font = body
            s, e = pdate(t["start"]), pdate(t.get("end", t["start"]))
            for off in range((e - s).days + 1):
                d = s + datetime.timedelta(days=off)
                col = DATE_COL_START + (d - start).days
                if DATE_COL_START <= col <= last_col:
                    ws.cell(row, col).fill = band
            row += 1
        row += 1  # spacer between phases


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", default=BUNDLED)
    ap.add_argument("--config", required=True)
    ap.add_argument("--output", required=True)
    args = ap.parse_args()
    if not os.path.exists(args.template):
        sys.exit(f"Template not found: {args.template}. Supply one with --template.")
    cfg = json.load(open(args.config))
    wb = openpyxl.load_workbook(args.template)
    # keep only the Project Schedule sheet
    for name in list(wb.sheetnames):
        if name != SCHEDULE_SHEET:
            wb.remove(wb[name])
    if SCHEDULE_SHEET not in wb.sheetnames:
        sys.exit(f"'{SCHEDULE_SHEET}' sheet not found in template.")
    build(wb[SCHEDULE_SHEET], cfg)
    wb.save(args.output)
    print(f"Saved: {args.output}")


if __name__ == "__main__":
    main()
